"""Script de investigación TEMPORAL (segunda ronda): la primera ronda confirmó que
el .xlsx de holdings de SPDR/State Street es un fichero real y parseable (firma ZIP
válida), a diferencia de iShares (devuelve el HTML del sitio, no el CSV) y VanEck
(HTML, sin confirmar si las posiciones están en el HTML o requieren JS).

Esta ronda comprueba DOS cosas con datos reales antes de proponer nada:
1) ¿El .xlsx trae columnas de sub-industria (GICS) o solo nombre/ticker/peso?
   Si trae sub-industria, se puede derivar SUBSECTOR_EXAMPLES real (no de memoria)
   agrupando las posiciones de XLF por su sub-industria (Banca/Seguros/Bróker...).
2) ¿Funciona igual para el ETF de semiconductores (XSD, familia SPDR) y para los
   tres sectores sin desglose de Ken French (XLU, XLB, XLC)?

Se borra en cuanto se ha capturado el resultado.
"""
import io
import requests

HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                         "(KHTML, like Gecko) Chrome/120.0 Safari/537.36"}

URL = ("https://www.ssga.com/us/en/individual/library-content/products/fund-data/"
       "etfs/us/holdings-daily-us-en-{ticker}.xlsx")


def inspect(ticker):
    url = URL.format(ticker=ticker.lower())
    try:
        r = requests.get(url, timeout=20, headers=HEADERS)
    except Exception as exc:  # noqa: BLE001
        print(f"=== {ticker}: ERROR {exc} ===")
        return
    print(f"=== {ticker}: HTTP {r.status_code}, {len(r.content)} bytes ===")
    if r.status_code != 200:
        return
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            print(f"  hoja: {sheet!r}  filas={ws.max_row}  cols={ws.max_column}")
        ws = wb[wb.sheetnames[0]]
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True)):
            print(f"    fila {i}: {row}")
    except Exception as exc:  # noqa: BLE001
        print(f"  ERROR al parsear con openpyxl: {exc}")


for t in ["XLF", "XLU", "XLB", "XLC", "XSD"]:
    inspect(t)
